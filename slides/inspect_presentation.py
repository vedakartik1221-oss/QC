from pptx import Presentation

prs = Presentation('slides/QuantumScreen_Presentation.pptx')
print(f"Total slides: {len(prs.slides)}")

for idx, slide in enumerate(prs.slides):
    shapes_info = []
    for shape in slide.shapes:
        type_str = str(shape.shape_type)
        text = shape.text_frame.text.strip().replace('\n', ' ') if shape.has_text_frame else ''
        if shape.has_chart:
            type_str = f"CHART: {shape.chart.chart_title.text_frame.text if shape.chart.has_title else 'No Title'}"
        shapes_info.append(f"[{type_str}] {text[:40]}")
    print(f"\n--- Slide {idx+1} ({len(slide.shapes)} shapes) ---")
    for s in shapes_info:
        print("  ", s.encode("ascii", "replace").decode())
